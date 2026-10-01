"""Iter-10 backend tests — KasirKu Pro Fase 1.

Scope:
  * GET /api/reports/summary|products|customers|operators|sales-items with
    period=custom + start/end YYYY-MM-DD (NEW)
  * Regression of same endpoints with period=daily|monthly|yearly
  * NEW GET /api/reports/export (xlsx multi-sheet; daily/monthly/custom)
  * Regression: POST /api/transactions still creates a sale
  * Regression: GET /api/products, /api/customers list still OK

Preview URL comes from EXPO_BACKEND_URL (set via conftest fallback).
"""
from __future__ import annotations

import io
import os
import time
import uuid
from datetime import datetime, timedelta, timezone

import pytest
import requests

BASE_URL = os.environ.get("EXPO_BACKEND_URL", "").rstrip("/")
API = f"{BASE_URL}/api"

if not BASE_URL:
    raise RuntimeError("EXPO_BACKEND_URL must be set for iter-10 tests")


# ---------- Fixtures ----------
@pytest.fixture(scope="module")
def owner_session():
    """Fresh TEST_Iter10 owner so date ranges are deterministic."""
    s = requests.Session()
    s.headers["Content-Type"] = "application/json"
    uname = f"iter10_{uuid.uuid4().hex[:8]}"
    r = s.post(f"{API}/auth/register", json={
        "business_name": "TEST_Iter10 Toko",
        "owner_name": "TEST_Iter10 Owner",
        "username": uname,
        "password": "password123",
    }, timeout=30)
    assert r.status_code == 200, f"register failed: {r.status_code} {r.text}"
    tok = r.json()["token"]
    s.headers["Authorization"] = f"Bearer {tok}"
    s.username = uname  # type: ignore[attr-defined]
    return s


@pytest.fixture(scope="module")
def seeded_sale(owner_session):
    """Create 1 product + 1 cash sale so reports have data for today."""
    s = owner_session
    r = s.post(f"{API}/products", json={
        "name": "TEST_Iter10 Kopi",
        "sku": f"IT10-{uuid.uuid4().hex[:5].upper()}",
        "category": "Minuman",
        "product_type": "goods",
        "base_unit": "Pcs",
        "price": 15000,
        "cost": 7000,
        "stock": 50,
        "min_stock": 5,
    }, timeout=20)
    assert r.status_code == 200, r.text
    product = r.json()

    tx = s.post(f"{API}/transactions", json={
        "transaction_type": "sale",
        "payment_method": "cash",
        "payment_method_label": "Tunai",
        "items": [{
            "product_id": product["id"],
            "name": product["name"],
            "quantity": 2,
            "price": 15000,
            "cost": 7000,
            "discount": 1000,  # per-item discount for Fase 1 UI
            "unit": "Pcs",
            "unit_factor": 1.0,
            "product_type": "goods",
        }],
        "paid_amount": 29000,
        "customer_name": "TEST_Iter10 Pelanggan",
    }, timeout=20)
    assert tx.status_code == 200, tx.text
    body = tx.json()
    assert body["transaction_type"] == "sale"
    assert body["total"] == 29000  # (15000*2) - 1000 item discount
    return {"product": product, "transaction": body}


def _today_iso_date() -> str:
    return datetime.now(timezone.utc).strftime("%Y-%m-%d")


def _yesterday_iso_date() -> str:
    return (datetime.now(timezone.utc) - timedelta(days=1)).strftime("%Y-%m-%d")


def _tomorrow_iso_date() -> str:
    return (datetime.now(timezone.utc) + timedelta(days=1)).strftime("%Y-%m-%d")


# ---------- /reports/summary ----------
class TestReportSummary:
    """Summary endpoint across periods incl. custom."""

    def test_summary_daily_regression(self, owner_session, seeded_sale):
        r = owner_session.get(f"{API}/reports/summary?period=daily", timeout=20)
        assert r.status_code == 200
        body = r.json()
        assert body["period"] == "daily"
        assert body["sales_total"] >= 29000
        assert body["transaction_count"] >= 1
        assert body["cash_in"] >= 29000

    def test_summary_monthly_regression(self, owner_session, seeded_sale):
        r = owner_session.get(f"{API}/reports/summary?period=monthly", timeout=20)
        assert r.status_code == 200
        assert r.json()["period"] == "monthly"
        assert r.json()["sales_total"] >= 29000

    def test_summary_yearly_regression(self, owner_session, seeded_sale):
        r = owner_session.get(f"{API}/reports/summary?period=yearly", timeout=20)
        assert r.status_code == 200
        assert r.json()["period"] == "yearly"

    def test_summary_custom_includes_today(self, owner_session, seeded_sale):
        s = _yesterday_iso_date()
        e = _tomorrow_iso_date()
        r = owner_session.get(f"{API}/reports/summary?period=custom&start={s}&end={e}", timeout=20)
        assert r.status_code == 200, r.text
        body = r.json()
        assert body["period"] == "custom"
        # today's seeded sale must be inside yesterday..tomorrow
        assert body["sales_total"] >= 29000
        assert body["transaction_count"] >= 1
        assert body["cash_in"] >= 29000

    def test_summary_custom_future_range_empty(self, owner_session, seeded_sale):
        future = (datetime.now(timezone.utc) + timedelta(days=10)).strftime("%Y-%m-%d")
        far = (datetime.now(timezone.utc) + timedelta(days=20)).strftime("%Y-%m-%d")
        r = owner_session.get(f"{API}/reports/summary?period=custom&start={future}&end={far}", timeout=20)
        assert r.status_code == 200
        body = r.json()
        assert body["period"] == "custom"
        assert body["sales_total"] == 0
        assert body["transaction_count"] == 0


# ---------- /reports/products|customers|operators|sales-items ----------
class TestDetailedReportsCustomRange:
    def test_products_custom_range(self, owner_session, seeded_sale):
        s, e = _yesterday_iso_date(), _tomorrow_iso_date()
        r = owner_session.get(f"{API}/reports/products?period=custom&start={s}&end={e}", timeout=20)
        assert r.status_code == 200, r.text
        body = r.json()
        assert body["period"] == "custom"
        assert isinstance(body["rows"], list)
        assert len(body["rows"]) >= 1
        row = next((x for x in body["rows"] if x["name"] == seeded_sale["product"]["name"]), None)
        assert row is not None, "seeded product not in rows"
        assert row["quantity"] == 2
        assert row["revenue"] == 30000  # price*qty (no discount subtracted here per impl)

    def test_products_daily_regression(self, owner_session, seeded_sale):
        r = owner_session.get(f"{API}/reports/products?period=daily", timeout=20)
        assert r.status_code == 200
        assert r.json()["period"] == "daily"

    def test_customers_custom_range(self, owner_session, seeded_sale):
        s, e = _yesterday_iso_date(), _tomorrow_iso_date()
        r = owner_session.get(f"{API}/reports/customers?period=custom&start={s}&end={e}", timeout=20)
        assert r.status_code == 200, r.text
        body = r.json()
        assert body["period"] == "custom"
        assert any(row["customer_name"] == "TEST_Iter10 Pelanggan" for row in body["rows"])

    def test_customers_monthly_regression(self, owner_session, seeded_sale):
        r = owner_session.get(f"{API}/reports/customers?period=monthly", timeout=20)
        assert r.status_code == 200
        assert r.json()["period"] == "monthly"

    def test_operators_custom_range(self, owner_session, seeded_sale):
        s, e = _yesterday_iso_date(), _tomorrow_iso_date()
        r = owner_session.get(f"{API}/reports/operators?period=custom&start={s}&end={e}", timeout=20)
        assert r.status_code == 200, r.text
        body = r.json()
        assert body["period"] == "custom"
        assert len(body["rows"]) >= 1
        # owner is the operator for seeded sale
        assert any(row["total_sales"] >= 29000 for row in body["rows"])

    def test_operators_yearly_regression(self, owner_session, seeded_sale):
        r = owner_session.get(f"{API}/reports/operators?period=yearly", timeout=20)
        assert r.status_code == 200
        assert r.json()["period"] == "yearly"

    def test_sales_items_custom_range(self, owner_session, seeded_sale):
        s, e = _yesterday_iso_date(), _tomorrow_iso_date()
        r = owner_session.get(f"{API}/reports/sales-items?period=custom&start={s}&end={e}", timeout=20)
        assert r.status_code == 200, r.text
        body = r.json()
        assert body["period"] == "custom"
        assert body["qty"] >= 2
        assert any(row["name"] == seeded_sale["product"]["name"] for row in body["rows"])

    def test_sales_items_future_range_empty(self, owner_session, seeded_sale):
        future = (datetime.now(timezone.utc) + timedelta(days=30)).strftime("%Y-%m-%d")
        far = (datetime.now(timezone.utc) + timedelta(days=40)).strftime("%Y-%m-%d")
        r = owner_session.get(f"{API}/reports/sales-items?period=custom&start={future}&end={far}", timeout=20)
        assert r.status_code == 200
        body = r.json()
        assert body["rows"] == []
        assert body["qty"] == 0
        assert body["total"] == 0


# ---------- /reports/export (NEW xlsx multi-sheet) ----------
class TestReportsExport:
    XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    EXPECTED_SHEETS = {"Ringkasan", "Per Produk", "Per Pelanggan", "Per Operator", "Detail Barang", "Detail Transaksi"}

    def _parse(self, blob: bytes):
        import openpyxl  # provided by pandas excel deps
        wb = openpyxl.load_workbook(io.BytesIO(blob), read_only=True)
        return wb

    def test_export_monthly_xlsx(self, owner_session, seeded_sale):
        r = owner_session.get(f"{API}/reports/export?period=monthly", timeout=60)
        assert r.status_code == 200, r.text[:300]
        assert r.headers.get("content-type", "").startswith(self.XLSX_MIME), r.headers.get("content-type")
        cd = r.headers.get("content-disposition", "")
        assert "attachment" in cd.lower() and ".xlsx" in cd.lower()
        assert len(r.content) > 500  # non-trivial xlsx
        wb = self._parse(r.content)
        sheets = set(wb.sheetnames)
        missing = self.EXPECTED_SHEETS - sheets
        assert not missing, f"missing sheets: {missing}; got {sheets}"
        # Ringkasan must contain the known metrics
        ws = wb["Ringkasan"]
        metrik_values = [row[0].value for row in ws.iter_rows(min_row=2) if row[0].value]
        assert "Total Penjualan" in metrik_values
        assert "Jumlah Transaksi Penjualan" in metrik_values
        assert "Saldo Kas" in metrik_values

    def test_export_custom_range_xlsx(self, owner_session, seeded_sale):
        s, e = _yesterday_iso_date(), _tomorrow_iso_date()
        r = owner_session.get(f"{API}/reports/export?period=custom&start={s}&end={e}", timeout=60)
        assert r.status_code == 200, r.text[:300]
        assert r.headers.get("content-type", "").startswith(self.XLSX_MIME)
        wb = self._parse(r.content)
        assert self.EXPECTED_SHEETS.issubset(set(wb.sheetnames))
        # Detail Transaksi should include the seeded sale
        ws = wb["Detail Transaksi"]
        rows = list(ws.iter_rows(values_only=True))
        assert len(rows) >= 2, "header + at least 1 tx expected"
        header = rows[0]
        assert "Invoice" in header and "Total" in header
        inv_idx = header.index("Invoice")
        invoices = [r[inv_idx] for r in rows[1:]]
        assert seeded_sale["transaction"]["invoice_no"] in invoices

    def test_export_custom_future_range_still_200(self, owner_session, seeded_sale):
        """Should return a valid xlsx with empty data sheets when range has no data."""
        future = (datetime.now(timezone.utc) + timedelta(days=30)).strftime("%Y-%m-%d")
        far = (datetime.now(timezone.utc) + timedelta(days=40)).strftime("%Y-%m-%d")
        r = owner_session.get(f"{API}/reports/export?period=custom&start={future}&end={far}", timeout=60)
        assert r.status_code == 200
        assert r.headers.get("content-type", "").startswith(self.XLSX_MIME)
        wb = self._parse(r.content)
        # All 6 sheets still exist even with no data
        assert self.EXPECTED_SHEETS.issubset(set(wb.sheetnames))

    def test_export_requires_auth(self):
        r = requests.get(f"{API}/reports/export?period=monthly", timeout=20)
        assert r.status_code == 401

    def test_export_daily_regression(self, owner_session, seeded_sale):
        r = owner_session.get(f"{API}/reports/export?period=daily", timeout=60)
        assert r.status_code == 200
        assert r.headers.get("content-type", "").startswith(self.XLSX_MIME)


# ---------- Regression: core endpoints ----------
class TestCoreRegression:
    def test_products_list_still_works(self, owner_session, seeded_sale):
        r = owner_session.get(f"{API}/products", timeout=20)
        assert r.status_code == 200
        rows = r.json()
        assert any(p["id"] == seeded_sale["product"]["id"] for p in rows)
        for p in rows:
            assert "_id" not in p
            assert "business_id" not in p

    def test_customers_list_still_works(self, owner_session, seeded_sale):
        r = owner_session.get(f"{API}/customers", timeout=20)
        assert r.status_code == 200
        assert isinstance(r.json(), list)

    def test_sales_list_still_works(self, owner_session, seeded_sale):
        r = owner_session.get(f"{API}/transactions?kind=sale", timeout=20)
        assert r.status_code == 200
        rows = r.json()
        assert any(tx["id"] == seeded_sale["transaction"]["id"] for tx in rows)

    def test_create_second_sale_regression(self, owner_session, seeded_sale):
        """POST /api/transactions still works end-to-end after Fase-1 additions."""
        s = owner_session
        pid = seeded_sale["product"]["id"]
        r = s.post(f"{API}/transactions", json={
            "transaction_type": "sale",
            "payment_method": "cash",
            "payment_method_label": "Tunai",
            "items": [{
                "product_id": pid,
                "name": "TEST_Iter10 Kopi",
                "quantity": 1,
                "price": 15000,
                "cost": 7000,
                "discount": 0,
                "unit": "Pcs",
                "unit_factor": 1.0,
                "product_type": "goods",
            }],
            "paid_amount": 15000,
            "customer_name": "Walk-in",
        }, timeout=20)
        assert r.status_code == 200, r.text
        tx = r.json()
        assert tx["total"] == 15000
        assert tx["outstanding"] == 0
        assert tx["invoice_no"].startswith("SA-")

    def test_bad_custom_dates_fallback_gracefully(self, owner_session, seeded_sale):
        """_parse_date returns None on garbage → impl falls back to today's start. Must not 500."""
        r = owner_session.get(
            f"{API}/reports/summary?period=custom&start=not-a-date&end=also-bad",
            timeout=20,
        )
        assert r.status_code == 200
        assert r.json()["period"] == "custom"
